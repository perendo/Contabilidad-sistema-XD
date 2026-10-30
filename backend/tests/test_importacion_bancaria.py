import io
from datetime import date
from decimal import Decimal

import openpyxl

from app.models.cuenta import Cuenta
from app.models.ejercicio import Ejercicio
from app.models.entidad import Entidad
from app.models.usuario import RolEnum


def _contexto(db, cif="A11111111", nombre="Entidad1", anio=2026):
    entidad = Entidad(cif=cif, nombre=nombre, activa=True)
    db.add(entidad)
    db.flush()
    ejercicio = Ejercicio(
        entidad_id=entidad.id,
        anio=anio,
        fecha_inicio=date(anio, 1, 1),
        fecha_fin=date(anio, 12, 31),
    )
    db.add(ejercicio)
    banco = Cuenta(codigo="572", nombre="Bancos c/c vista", grupo="5", activa=True)
    puente = Cuenta(codigo="555", nombre="Partidas pendientes de aplicación", grupo="5", activa=True)
    db.add_all([banco, puente])
    db.commit()
    return entidad, ejercicio, banco, puente


def _excel_movimientos(filas):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja["A8"] = "Fecha Operación"
    hoja["B8"] = "Fecha Valor"
    hoja["C8"] = "Concepto"
    hoja["D8"] = "Importe"
    hoja["E8"] = "Divisa"
    for indice, (fecha, concepto, importe) in enumerate(filas, start=9):
        hoja.cell(row=indice, column=1, value=fecha)
        hoja.cell(row=indice, column=3, value=concepto)
        hoja.cell(row=indice, column=4, value=importe)
        hoja.cell(row=indice, column=5, value="EUR")
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


def _login(client, crear_usuario, email="tesorero@rendo.com"):
    crear_usuario(email=email, rol=RolEnum.TESORERO)
    login = client.post("/api/v1/auth/login", json={"email": email, "password": "Admin67@"})
    client.headers.update({"Authorization": f"Bearer {login.json()['access_token']}"})
    return client


def test_importar_excel_valido_crea_un_borrador_por_movimiento(client, crear_usuario, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    _login(client, crear_usuario)
    archivo = _excel_movimientos(
        [
            ("15/01/2026", "Ingreso donativo", 100.00),
            ("20/01/2026", "Pago proveedor", -50.00),
            ("25/01/2026", "Transferencia", 25.00),
        ]
    )

    respuesta = client.post(
        "/api/v1/asientos/importar",
        data={
            "entidad_id": str(entidad.id),
            "ejercicio_id": str(ejercicio.id),
            "cuenta_id": str(banco.id),
        },
        files={"archivo": ("movimientos.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert datos["creados"] == 3
    assert datos["omitidos"] == 0

    listado = client.get(
        "/api/v1/asientos/",
        params={"entidad_id": str(entidad.id), "ejercicio_id": str(ejercicio.id)},
    )
    assert listado.status_code == 200
    borradores = [a for a in listado.json() if a.get("estado") == "borrador"]
    assert len(borradores) == 3


def test_importar_movimiento_positivo_genera_borrador_con_banco_al_debe(client, crear_usuario, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    _login(client, crear_usuario)
    archivo = _excel_movimientos([("15/01/2026", "Ingreso donativo", 100.00)])

    respuesta = client.post(
        "/api/v1/asientos/importar",
        data={
            "entidad_id": str(entidad.id),
            "ejercicio_id": str(ejercicio.id),
            "cuenta_id": str(banco.id),
        },
        files={"archivo": ("movimientos.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert respuesta.status_code == 200
    assert respuesta.json()["creados"] == 1

    listado = client.get(
        "/api/v1/asientos/",
        params={"entidad_id": str(entidad.id), "ejercicio_id": str(ejercicio.id)},
    )
    borrador = next(a for a in listado.json() if a["estado"] == "borrador")
    apuntes = borrador["apuntes"]
    assert len(apuntes) == 2
    debe = next(a for a in apuntes if a["tipo"] == "debe")
    haber = next(a for a in apuntes if a["tipo"] == "haber")
    assert debe["cuenta_id"] == str(banco.id)
    assert haber["cuenta_id"] == str(puente.id)
def test_importar_movimiento_negativo_genera_borrador_con_banco_al_haber(client, crear_usuario, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    _login(client, crear_usuario)
    archivo = _excel_movimientos([("20/01/2026", "Pago proveedor", -50.00)])

    respuesta = client.post(
        "/api/v1/asientos/importar",
        data={
            "entidad_id": str(entidad.id),
            "ejercicio_id": str(ejercicio.id),
            "cuenta_id": str(banco.id),
        },
        files={"archivo": ("movimientos.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert respuesta.status_code == 200
    assert respuesta.json()["creados"] == 1

    listado = client.get(
        "/api/v1/asientos/",
        params={"entidad_id": str(entidad.id), "ejercicio_id": str(ejercicio.id)},
    )
    borrador = next(a for a in listado.json() if a["estado"] == "borrador")
    apuntes = borrador["apuntes"]
    assert len(apuntes) == 2
    debe = next(a for a in apuntes if a["tipo"] == "debe")
    haber = next(a for a in apuntes if a["tipo"] == "haber")
    assert debe["cuenta_id"] == str(puente.id)
    assert haber["cuenta_id"] == str(banco.id)
def test_importar_archivo_invalido_no_crea_ningun_borrador(client, crear_usuario, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    _login(client, crear_usuario)

    respuesta = client.post(
        "/api/v1/asientos/importar",
        data={
            "entidad_id": str(entidad.id),
            "ejercicio_id": str(ejercicio.id),
            "cuenta_id": str(banco.id),
        },
        files={"archivo": ("movimientos.txt", io.BytesIO(b"esto no es un excel"), "text/plain")},
    )

    assert respuesta.status_code == 400

    listado = client.get(
        "/api/v1/asientos/",
        params={"entidad_id": str(entidad.id), "ejercicio_id": str(ejercicio.id)},
    )
def test_importar_con_fecha_fuera_del_ejercicio_rechaza_todo(client, crear_usuario, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    _login(client, crear_usuario)
    archivo = _excel_movimientos(
        [
            ("15/01/2026", "Ingreso donativo", 100.00),
            ("15/01/2027", "Movimiento fuera de ejercicio", 10.00),
        ]
    )

    respuesta = client.post(
        "/api/v1/asientos/importar",
        data={
            "entidad_id": str(entidad.id),
            "ejercicio_id": str(ejercicio.id),
            "cuenta_id": str(banco.id),
        },
        files={"archivo": ("movimientos.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert respuesta.status_code == 400

    listado = client.get(
        "/api/v1/asientos/",
        params={"entidad_id": str(entidad.id), "ejercicio_id": str(ejercicio.id)},
    )
def test_importar_movimiento_duplicado_se_omite_sin_duplicar(client, crear_usuario, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    _login(client, crear_usuario)
    archivo = _excel_movimientos(
        [
            ("15/01/2026", "Ingreso donativo", 100.00),
            ("20/01/2026", "Pago proveedor", -50.00),
            ("25/01/2026", "Transferencia", 25.00),
        ]
    )

    primera = client.post(
        "/api/v1/asientos/importar",
        data={
            "entidad_id": str(entidad.id),
            "ejercicio_id": str(ejercicio.id),
            "cuenta_id": str(banco.id),
        },
        files={"archivo": ("movimientos.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert primera.status_code == 200
    assert primera.json()["creados"] == 3

    segunda = client.post(
        "/api/v1/asientos/importar",
        data={
            "entidad_id": str(entidad.id),
            "ejercicio_id": str(ejercicio.id),
            "cuenta_id": str(banco.id),
        },
        files={"archivo": ("movimientos.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert segunda.status_code == 200
    assert segunda.json()["creados"] == 0
    assert segunda.json()["omitidos"] == 3

    listado = client.get(
        "/api/v1/asientos/",
        params={"entidad_id": str(entidad.id), "ejercicio_id": str(ejercicio.id)},
    )
    borradores = [a for a in listado.json() if a["estado"] == "borrador"]
    assert len(borradores) == 3