import io
from datetime import date

import openpyxl

from app.models.cuenta import Cuenta
from app.models.ejercicio import Ejercicio
from app.models.entidad import Entidad


def _contexto(db, cif="B22222222", nombre="Entidad2", anio=2026):
    entidad = Entidad(cif=cif, nombre=nombre, activa=True)
    db.add(entidad)
    db.flush()
    ejercicio = Ejercicio(
        entidad_id=entidad.id,
        anio=anio,
        fecha_inicio=date(anio, 1, 1),
        fecha_fin=date(anio, 12, 31),
    )
    db.add(ejercicio)
    banco = Cuenta(codigo="572", nombre="Bancos c/c vista", grupo="5", activa=True)
    puente = Cuenta(codigo="555", nombre="Partidas pendientes de aplicación", grupo="5", activa=True)
    db.add_all([banco, puente])
    db.commit()
    return entidad, ejercicio, banco, puente


def _excel_movimientos(filas):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja["A8"] = "Fecha Operación"
    hoja["C8"] = "Concepto"
    hoja["D8"] = "Importe"
    for indice, (fecha, concepto, importe) in enumerate(filas, start=9):
        hoja.cell(row=indice, column=1, value=fecha)
        hoja.cell(row=indice, column=3, value=concepto)
        hoja.cell(row=indice, column=4, value=importe)
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


def _importar(cliente, entidad, ejercicio, banco):
    archivo = _excel_movimientos([("15/01/2026", "Ingreso donativo", 100.00)])
    respuesta = cliente.post(
        "/api/v1/asientos/importar",
        data={
            "entidad_id": str(entidad.id),
            "ejercicio_id": str(ejercicio.id),
            "cuenta_id": str(banco.id),
        },
        files={"archivo": ("movimientos.xlsx", archivo, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta.status_code == 200
    listado = cliente.get(
        "/api/v1/asientos/",
        params={"entidad_id": str(entidad.id), "ejercicio_id": str(ejercicio.id)},
    )
    return next(a for a in listado.json() if a["estado"] == "borrador")


def test_editar_borrador_actualiza_concepto_y_apuntes(cliente_tesorero, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    borrador = _importar(cliente_tesorero, entidad, ejercicio, banco)

    respuesta = cliente_tesorero.patch(
        f"/api/v1/asientos/{borrador['id']}",
        json={
            "concepto": "Donativo revisado",
            "apuntes": [
                {"cuenta_id": str(banco.id), "tipo": "debe", "importe": "120.00"},
                {"cuenta_id": str(puente.id), "tipo": "haber", "importe": "120.00"},
            ],
        },
    )

    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert datos["concepto"] == "Donativo revisado"
    assert datos["estado"] == "borrador"
    assert len(datos["apuntes"]) == 2
    assert datos["apuntes"][0]["importe"] in ("120.00", "120.0")


def test_aprobar_borrador_lo_deja_registrado(cliente_tesorero, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    borrador = _importar(cliente_tesorero, entidad, ejercicio, banco)

    respuesta = cliente_tesorero.post(f"/api/v1/asientos/{borrador['id']}/aprobar")

    assert respuesta.status_code == 200
    assert respuesta.json()["estado"] == "registrado"


def test_rechazar_borrador_lo_elimina(cliente_tesorero, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    borrador = _importar(cliente_tesorero, entidad, ejercicio, banco)

    respuesta = cliente_tesorero.delete(f"/api/v1/asientos/{borrador['id']}")

    assert respuesta.status_code == 204

    listado = cliente_tesorero.get(
        "/api/v1/asientos/",
        params={"entidad_id": str(entidad.id), "ejercicio_id": str(ejercicio.id)},
    )
    assert all(a["id"] != borrador["id"] for a in listado.json())


def test_no_se_puede_editar_un_asiento_registrado(cliente_tesorero, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    borrador = _importar(cliente_tesorero, entidad, ejercicio, banco)
    cliente_tesorero.post(f"/api/v1/asientos/{borrador['id']}/aprobar")

    respuesta = cliente_tesorero.patch(
        f"/api/v1/asientos/{borrador['id']}",
        json={"concepto": "Intento de edición"},
    )

    assert respuesta.status_code == 400


def test_no_se_puede_aprobar_un_asiento_registrado(cliente_tesorero, db):
    entidad, ejercicio, banco, puente = _contexto(db)
    borrador = _importar(cliente_tesorero, entidad, ejercicio, banco)
    cliente_tesorero.post(f"/api/v1/asientos/{borrador['id']}/aprobar")

    respuesta = cliente_tesorero.post(f"/api/v1/asientos/{borrador['id']}/aprobar")

    assert respuesta.status_code == 400


def test_solo_lectura_no_puede_editar_borrador(cliente_solo_lectura, db):
    from app.models.usuario import RolEnum
    # crear contexto con un usuario tesorero no es necesario; usamos cliente solo lectura directamente
    respuesta = cliente_solo_lectura.patch(
        f"/api/v1/asientos/00000000-0000-0000-0000-000000000000",
        json={"concepto": "x"},
    )
    assert respuesta.status_code == 403