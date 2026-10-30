import io
import logging
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Annotated
from uuid import UUID, uuid4

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import FileResponse
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from app.api.deps import get_db, solo_tesorero_o_admin, usuario_actual
from app.core.config import BASE_DIR, settings
from app.models.asiento import Apunte, Asiento, DocumentoAdjunto, EstadoAsiento, TipoApunte
from app.models.cuenta import Cuenta
from app.models.ejercicio import Ejercicio
from app.models.entidad import Entidad
from app.models.usuario import Usuario
from app.schemas.asiento import AsientoCreate, AsientoRead, AsientoUpdate, DocumentoRead

router = APIRouter()
logger = logging.getLogger(__name__)

MAX_TAMANO_IMPORTACION = 10 * 1024 * 1024
MAX_FILAS_IMPORTACION = 5000


@router.post("/", response_model=AsientoRead, status_code=status.HTTP_201_CREATED)
def crear_asiento(
    asiento_in: AsientoCreate,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(solo_tesorero_o_admin),
):

    # Obtener siguiente número de asiento por entidad+ejercicio
    max_num = (
        db.query(func.max(Asiento.numero_asiento))
        .filter(Asiento.entidad_id == asiento_in.entidad_id, Asiento.ejercicio_id == asiento_in.ejercicio_id)
        .scalar()
        or 0
    )
    nuevo_numero = max_num + 1

    db_asiento = Asiento(
        entidad_id=asiento_in.entidad_id,
        ejercicio_id=asiento_in.ejercicio_id,
        numero_asiento=nuevo_numero,
        fecha=asiento_in.fecha,
        concepto=asiento_in.concepto,
        descripcion=asiento_in.descripcion,
        registrado_por_id=usuario.id,
    )

    db.add(db_asiento)
    db.flush()

    for apunte_in in asiento_in.apuntes:
        db_apunte = Apunte(
            asiento_id=db_asiento.id,
            cuenta_id=apunte_in.cuenta_id,
            tipo=apunte_in.tipo,
            importe=apunte_in.importe,
        )
        db.add(db_apunte)

    db.commit()
    db.refresh(db_asiento)
    return db_asiento


@router.patch("/{asiento_id}", response_model=AsientoRead)
def actualizar_asiento(
    asiento_id: UUID,
    asiento_in: AsientoUpdate,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(solo_tesorero_o_admin),
):
    asiento = (
        db.query(Asiento)
        .options(selectinload(Asiento.apuntes), selectinload(Asiento.ejercicio))
        .filter(Asiento.id == asiento_id)
        .first()
    )
    if asiento is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Asiento no encontrado")
    if asiento.estado != EstadoAsiento.BORRADOR:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Solo se pueden modificar borradores",
        )

    if asiento_in.fecha is not None:
        if asiento_in.fecha < asiento.ejercicio.fecha_inicio or asiento_in.fecha > asiento.ejercicio.fecha_fin:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="La fecha está fuera del ejercicio seleccionado",
            )
        asiento.fecha = asiento_in.fecha
    if asiento_in.concepto is not None:
        concepto_limpio = asiento_in.concepto.strip()
        if not concepto_limpio:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El concepto no puede estar vacío",
            )
        asiento.concepto = concepto_limpio
    if asiento_in.descripcion is not None:
        asiento.descripcion = asiento_in.descripcion

    if asiento_in.apuntes is not None:
        for apunte_in in asiento_in.apuntes:
            cuenta = db.query(Cuenta).filter(Cuenta.id == apunte_in.cuenta_id).first()
            if cuenta is None or not cuenta.activa or cuenta.entidad_id not in (None, asiento.entidad_id):
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Cuenta de apunte no válida",
                )
        for apunte in list(asiento.apuntes):
            db.delete(apunte)
        db.flush()
        for apunte_in in asiento_in.apuntes:
            db.add(
                Apunte(
                    asiento_id=asiento.id,
                    cuenta_id=apunte_in.cuenta_id,
                    tipo=apunte_in.tipo,
                    importe=apunte_in.importe,
                )
            )

    asiento.modificado_en = datetime.utcnow()
    db.commit()
    asiento = (
        db.query(Asiento)
        .options(selectinload(Asiento.apuntes), selectinload(Asiento.documentos))
        .filter(Asiento.id == asiento_id)
        .first()
    )
    return asiento


@router.post("/{asiento_id}/aprobar", response_model=AsientoRead)
def aprobar_asiento(
    asiento_id: UUID,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(solo_tesorero_o_admin),
):
    asiento = db.query(Asiento).filter(Asiento.id == asiento_id).first()
    if asiento is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Asiento no encontrado")
    if asiento.estado != EstadoAsiento.BORRADOR:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Solo se pueden aprobar borradores",
        )
    asiento.estado = EstadoAsiento.REGISTRADO
    db.commit()
    asiento = (
        db.query(Asiento)
        .options(selectinload(Asiento.apuntes), selectinload(Asiento.documentos))
        .filter(Asiento.id == asiento_id)
        .first()
    )
    return asiento


@router.delete("/{asiento_id}", status_code=status.HTTP_204_NO_CONTENT)
def rechazar_asiento(
    asiento_id: UUID,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(solo_tesorero_o_admin),
):
    asiento = db.query(Asiento).filter(Asiento.id == asiento_id).first()
    if asiento is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Asiento no encontrado")
    if asiento.estado != EstadoAsiento.BORRADOR:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Solo se pueden rechazar borradores",
        )
    documentos = db.query(DocumentoAdjunto).filter(DocumentoAdjunto.asiento_id == asiento_id).all()
    for documento in documentos:
        try:
            _resolver_documento(documento.ruta).unlink(missing_ok=True)
        except Exception:
            logger.exception("No se pudo borrar el archivo del documento %s", documento.id)
    db.delete(asiento)
    db.commit()
    return None


MAX_ARCHIVOS = 20
LIMITE_TAMANO = 10 * 1024 * 1024
LIMITE_TOTAL = 50 * 1024 * 1024
MAX_CLAVE_IDEMPOTENCIA = 128
MAX_NOMBRE_ARCHIVO = 255
DIRECTORIO_LEGADO = BASE_DIR / "documentos"
FIRMAS = (
    (b"%PDF-", "application/pdf", ".pdf"),
    (b"\xff\xd8\xff", "image/jpeg", ".jpg"),
    (b"\x89PNG\r\n\x1a\n", "image/png", ".png"),
)


def _nombre_carpeta(valor: str, entidad_id: UUID) -> str:
    """Nombre de carpeta seguro a partir del CIF, con respaldo si queda vacío."""
    limpio = "".join(caracter for caracter in valor if caracter.isalnum() or caracter in "-_")
    return limpio or f"entidad-{entidad_id}"


def _dentro_de(base: Path, destino: Path) -> bool:
    return destino == base or destino.is_relative_to(base)


def _directorio_documentos(asiento: Asiento) -> Path:
    base = settings.DIRECTORIO_DATOS.resolve()
    directorio = (
        base / _nombre_carpeta(asiento.entidad.cif, asiento.entidad_id) / str(asiento.ejercicio.anio)
    ).resolve()
    if not _dentro_de(base, directorio):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se pudo determinar la carpeta del documento",
        )
    return directorio


def _resolver_documento(ruta: str) -> Path:
    """Resuelve la ruta guardada sin permitir que escape de los directorios de datos."""
    base = settings.DIRECTORIO_DATOS.resolve()
    normalizada = ruta.replace("\\", "/")
    candidata = Path(normalizada)
    if candidata.is_absolute():
        destino = candidata.resolve()
        legado = DIRECTORIO_LEGADO.resolve()
        if not (_dentro_de(base, destino) or _dentro_de(legado, destino)):
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Documento no encontrado")
        return destino

    partes = [parte for parte in normalizada.split("/") if parte not in ("", ".")]
    destino = base.joinpath(*partes).resolve()
    if not _dentro_de(base, destino):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Documento no encontrado")
    return destino


def _nombre_archivo(archivo: UploadFile, por_defecto: str) -> str:
    nombre = re.split(r"[\\/]", archivo.filename or "")[-1]
    nombre = "".join(caracter for caracter in nombre if caracter.isprintable()).strip()
    return nombre[:MAX_NOMBRE_ARCHIVO] or por_defecto


def _leer_documento(archivo: UploadFile) -> tuple[bytes, str, str]:
    """Lee un documento y devuelve su contenido, tipo real y extensión."""
    if archivo.size is not None and archivo.size > LIMITE_TAMANO:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Cada archivo puede ocupar como máximo 10 MB",
        )

    contenido = archivo.file.read(LIMITE_TAMANO + 1)
    if not contenido:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="El archivo está vacío")
    if len(contenido) > LIMITE_TAMANO:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Cada archivo puede ocupar como máximo 10 MB",
        )

    for firma, tipo, extension in FIRMAS:
        if contenido.startswith(firma):
            return contenido, tipo, extension

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Solo se admiten archivos PDF, JPG o PNG",
    )


def _borrar_archivos(rutas: list[Path], directorio: Path) -> None:
    for ruta in rutas:
        ruta.unlink(missing_ok=True)
    try:
        directorio.rmdir()
    except OSError:
        pass


def _documentos_con_clave(db: Session, asiento_id: UUID, clave: str) -> list[DocumentoAdjunto]:
    return (
        db.query(DocumentoAdjunto)
        .filter(
            DocumentoAdjunto.asiento_id == asiento_id,
            DocumentoAdjunto.clave_idempotencia == clave,
        )
        .order_by(DocumentoAdjunto.creado_en, DocumentoAdjunto.id)
        .all()
    )


def _parsear_movimientos_excel(contenido: bytes) -> list[dict]:
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
    except Exception as error:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo no es un Excel válido",
        ) from error

    hoja = libro.active
    fila_cabecera = None
    for fila in hoja.iter_rows(min_row=1, max_row=20):
        for celda in fila:
            if isinstance(celda.value, str) and celda.value.strip().lower() in (
                "fecha operación",
                "fecha operacion",
            ):
                fila_cabecera = celda.row
                break
        if fila_cabecera is not None:
            break
    if fila_cabecera is None:
        fila_cabecera = 8

    movimientos: list[dict] = []
    for fila in hoja.iter_rows(min_row=fila_cabecera + 1):
        fecha_valor = fila[0].value if len(fila) > 0 else None
        concepto = fila[2].value if len(fila) > 2 else None
        importe = fila[3].value if len(fila) > 3 else None
        if fecha_valor is None and concepto is None and importe is None:
            continue
        if fecha_valor is None or concepto is None or importe is None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El archivo contiene filas incompletas",
            )

        if isinstance(fecha_valor, datetime):
            fecha = fecha_valor.date()
        elif isinstance(fecha_valor, date):
            fecha = fecha_valor
        elif isinstance(fecha_valor, str):
            try:
                fecha = datetime.strptime(fecha_valor.strip(), "%d/%m/%Y").date()
            except ValueError as error:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Fecha de movimiento no válida",
                ) from error
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Fecha de movimiento no válida",
            )

        try:
            importe_decimal = Decimal(str(importe))
        except Exception as error:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Importe de movimiento no válido",
            ) from error

        if not importe_decimal.is_finite():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Importe de movimiento no válido",
            )
        if importe_decimal.as_tuple().exponent < -2:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El importe no puede tener más de 2 decimales",
            )

        if importe_decimal == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El importe de un movimiento no puede ser cero",
            )

        concepto_limpio = str(concepto).strip()
        if not concepto_limpio:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="El concepto de un movimiento no puede estar vacío",
            )

        movimientos.append(
            {
                "fecha": fecha,
                "concepto": concepto_limpio,
                "importe": importe_decimal,
            }
        )

    if not movimientos:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo no contiene movimientos",
        )
    if len(movimientos) > MAX_FILAS_IMPORTACION:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo contiene demasiados movimientos",
        )
    return movimientos


def _cuenta_puente(db: Session, entidad_id: UUID) -> Cuenta:
    cuenta = (
        db.query(Cuenta)
        .filter(Cuenta.codigo == "555", Cuenta.entidad_id == entidad_id, Cuenta.activa == True)
        .first()
    )
    if cuenta is None:
        cuenta = (
            db.query(Cuenta)
            .filter(Cuenta.codigo == "555", Cuenta.entidad_id.is_(None), Cuenta.activa == True)
            .first()
        )
    if cuenta is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No existe la cuenta puente 555",
        )
    return cuenta


@router.post(
    "/{asiento_id}/documentos",
    response_model=list[DocumentoRead],
    status_code=status.HTTP_201_CREATED,
)
def subir_documentos(
    asiento_id: UUID,
    archivos: list[UploadFile] = File(...),
    clave_idempotencia: str | None = Form(None, max_length=MAX_CLAVE_IDEMPOTENCIA),
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(solo_tesorero_o_admin),
):
    asiento = db.query(Asiento).filter(Asiento.id == asiento_id).first()
    if not asiento:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Asiento no encontrado")

    if len(archivos) > MAX_ARCHIVOS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"No se pueden adjuntar más de {MAX_ARCHIVOS} documentos a la vez",
        )
    if sum(archivo.size or 0 for archivo in archivos) > LIMITE_TOTAL:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El conjunto de documentos supera los 50 MB",
        )
    if clave_idempotencia and len(archivos) != 1:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="La clave de idempotencia solo admite un archivo por petición",
        )

    # Se valida el contenido antes de mirar la clave: un reintento con un archivo
    # inválido nunca devuelve un 201 con los documentos antiguos.
    leidos = [(archivo, *_leer_documento(archivo)) for archivo in archivos]
    if sum(len(contenido) for _, contenido, _, _ in leidos) > LIMITE_TOTAL:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El conjunto de documentos supera los 50 MB",
        )

    if clave_idempotencia:
        existentes = _documentos_con_clave(db, asiento.id, clave_idempotencia)
        if existentes:
            return existentes

    directorio = _directorio_documentos(asiento)
    escritos: list[Path] = []
    documentos: list[DocumentoAdjunto] = []
    try:
        directorio.mkdir(parents=True, exist_ok=True)
        for archivo, contenido, tipo, extension in leidos:
            destino = directorio / f"{uuid4().hex}{extension}"
            escritos.append(destino)
            destino.write_bytes(contenido)
            documento = DocumentoAdjunto(
                asiento_id=asiento.id,
                nombre=_nombre_archivo(archivo, destino.name),
                tipo_contenido=tipo,
                tamano=len(contenido),
                ruta=destino.relative_to(settings.DIRECTORIO_DATOS.resolve()).as_posix(),
                clave_idempotencia=clave_idempotencia,
                subido_por_id=usuario.id,
            )
            db.add(documento)
            documentos.append(documento)
        db.commit()
    except IntegrityError as error:
        db.rollback()
        _borrar_archivos(escritos, directorio)
        if clave_idempotencia:
            existentes = _documentos_con_clave(db, asiento.id, clave_idempotencia)
            if existentes:
                return existentes
        logger.exception("Conflicto al guardar documentos del asiento %s", asiento.id)
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="No se pudo completar la subida; inténtalo de nuevo",
        ) from error
    except HTTPException:
        db.rollback()
        _borrar_archivos(escritos, directorio)
        raise
    except Exception as error:
        db.rollback()
        _borrar_archivos(escritos, directorio)
        logger.exception("No se pudieron guardar los documentos del asiento %s", asiento.id)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="No se pudieron guardar los documentos",
        ) from error

    for documento in documentos:
        db.refresh(documento)
    return documentos


@router.get("/{asiento_id}/documentos/{documento_id}/descarga")
def descargar_documento(
    asiento_id: UUID,
    documento_id: UUID,
    db: Session = Depends(get_db),
):
    documento = (
        db.query(DocumentoAdjunto)
        .filter(
            DocumentoAdjunto.id == documento_id,
            DocumentoAdjunto.asiento_id == asiento_id,
        )
        .first()
    )
    if not documento:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Documento no encontrado")

    ruta = _resolver_documento(documento.ruta)
    if not ruta.is_file():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Documento no encontrado")

    return FileResponse(
        ruta,
        media_type=documento.tipo_contenido,
        filename=documento.nombre,
    )


@router.post("/importar", status_code=status.HTTP_200_OK)
def importar_movimientos(
    archivo: UploadFile = File(...),
    entidad_id: UUID = Form(...),
    ejercicio_id: UUID = Form(...),
    cuenta_id: UUID = Form(...),
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(solo_tesorero_o_admin),
):
    entidad = db.query(Entidad).filter(Entidad.id == entidad_id).first()
    if entidad is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Entidad no encontrada")

    ejercicio = (
        db.query(Ejercicio)
        .filter(Ejercicio.id == ejercicio_id, Ejercicio.entidad_id == entidad_id)
        .first()
    )
    if ejercicio is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Ejercicio no encontrado")

    cuenta_banco = db.query(Cuenta).filter(Cuenta.id == cuenta_id).first()
    if cuenta_banco is None or not cuenta_banco.activa:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cuenta bancaria no válida")
    if cuenta_banco.entidad_id not in (None, entidad_id):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cuenta bancaria no válida")
    if not cuenta_banco.codigo.startswith("57"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="La cuenta seleccionada no es una cuenta bancaria")

    contenido = archivo.file.read(MAX_TAMANO_IMPORTACION + 1)
    if not contenido:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="El archivo está vacío")
    if len(contenido) > MAX_TAMANO_IMPORTACION:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo supera los 10 MB",
        )

    movimientos = _parsear_movimientos_excel(contenido)

    for movimiento in movimientos:
        if movimiento["fecha"] < ejercicio.fecha_inicio or movimiento["fecha"] > ejercicio.fecha_fin:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Hay movimientos con fecha fuera del ejercicio seleccionado",
            )

    puente = _cuenta_puente(db, entidad_id)
    if puente.id == cuenta_banco.id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="La cuenta puente no puede coincidir con la cuenta bancaria",
        )

    creados: list[Asiento] = []
    omitidos = 0
    max_num = (
        db.query(func.max(Asiento.numero_asiento))
        .filter(Asiento.entidad_id == entidad_id, Asiento.ejercicio_id == ejercicio_id)
        .scalar()
        or 0
    )

    for movimiento in movimientos:
        duplicado = (
            db.query(Apunte)
            .join(Asiento)
            .filter(
                Asiento.entidad_id == entidad_id,
                Asiento.ejercicio_id == ejercicio_id,
                Asiento.fecha == movimiento["fecha"],
                Asiento.concepto == movimiento["concepto"],
                Apunte.cuenta_id == cuenta_banco.id,
                Apunte.importe == abs(movimiento["importe"]),
            )
            .first()
        )
        if duplicado is not None:
            omitidos += 1
            continue

        importe = abs(movimiento["importe"])
        if movimiento["importe"] > 0:
            apuntes = [
                (cuenta_banco.id, TipoApunte.DEBE, importe),
                (puente.id, TipoApunte.HABER, importe),
            ]
        else:
            apuntes = [
                (puente.id, TipoApunte.DEBE, importe),
                (cuenta_banco.id, TipoApunte.HABER, importe),
            ]

        max_num += 1
        asiento = Asiento(
            entidad_id=entidad_id,
            ejercicio_id=ejercicio_id,
            numero_asiento=max_num,
            fecha=movimiento["fecha"],
            concepto=movimiento["concepto"],
            descripcion="Importado desde extracto bancario",
            estado=EstadoAsiento.BORRADOR,
            registrado_por_id=usuario.id,
        )
        db.add(asiento)
        db.flush()
        for cuenta_id_apunte, tipo, valor in apuntes:
            db.add(
                Apunte(
                    asiento_id=asiento.id,
                    cuenta_id=cuenta_id_apunte,
                    tipo=tipo,
                    importe=valor,
                )
            )
        creados.append(asiento)

    db.commit()
    for asiento in creados:
        db.refresh(asiento)

    return {
        "creados": len(creados),
        "omitidos": omitidos,
        "errores": [],
    }


@router.get("/", response_model=list[AsientoRead])
def listar_asientos(entidad_id: UUID, ejercicio_id: UUID, skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    asientos = (
        db.query(Asiento)
        .options(selectinload(Asiento.apuntes), selectinload(Asiento.documentos))
        .filter(Asiento.entidad_id == entidad_id, Asiento.ejercicio_id == ejercicio_id)
        .order_by(Asiento.fecha.desc(), Asiento.numero_asiento.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return asientos



@router.get("/{asiento_id}", response_model=AsientoRead)
def obtener_asiento(asiento_id: UUID, db: Session = Depends(get_db)):
    asiento = (
        db.query(Asiento)
        .options(selectinload(Asiento.apuntes), selectinload(Asiento.documentos))
        .filter(Asiento.id == asiento_id)
        .first()
    )
    if not asiento:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Asiento no encontrado")
    return asiento
